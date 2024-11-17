import pandas as pd
from pathlib import Path
from enum import Enum
from datetime import datetime
import numpy as np
from matplotlib import pyplot as plt
from sqlalchemy.types import Integer, NVARCHAR, REAL
from sqlalchemy import create_engine
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from dateutil.relativedelta import relativedelta

DATE_FORMAT = "%Y-%m-%d %H:%M:%S"


class TransactionType(Enum):
    """
    enum representing the two possible transaction types
    """
    EXPENSE = 0
    INCOME = 1


class Pynance:

    def __init__(self, db_file='sqlite:///transactions.db'):
        self.db_file = db_file
        self.transactions = self.load_transactions()

        self.dtypes = {
            "amount": REAL(),
            "created": NVARCHAR(length=255),
            "transaction_type": Integer(),
            "category": NVARCHAR(length=255),
            "description": NVARCHAR(length=255),
        }

    def load_transactions(self):
        if not Path(self.db_file).exists():
            return []
        else:
            return pd.read_sql_query(
                'SELECT * FROM transactions', self.db_file
            ).values.tolist()

    def save_transactions(self):
        df = pd.DataFrame(self.transactions)
        eng = create_engine(self.db_file)
        df.to_sql('transactions',
                  eng,
                  if_exists='replace',
                  index=False,
                  dtype=self.dtypes
                  )

    def process_transaction(self,
                            amt: float,
                            cat: str,
                            trans_type: TransactionType = TransactionType.EXPENSE,
                            desc: str = ""
                            ) -> None:
        """
        method to add transaction to database
        :param amt: amt of transaction
        :param cat: category of transaction
        :param trans_type: income or expense
        :param desc: description of transaction
        :return:
        """
        self.transactions.append({
            "amount": amt,
            "created": datetime.now().strftime(DATE_FORMAT),
            "transaction_type": trans_type.value,
            "category": cat.title(),
            "description": desc.lower(),
        })
        self.save_transactions()

    def get_summary_data(self):

        income = [trans['amount'] for trans in self.transactions if
                  trans['transaction_type'] == TransactionType.INCOME.value]
        expense = [trans['amount'] for trans in self.transactions if
                   trans['transaction_type'] == TransactionType.EXPENSE.value]

        return {
            "IncomeTotal": sum(income),
            "IncomeCount": len(income),
            "ExpenseTotal": sum(expense),
            "ExpenseCount": len(expense),
            "CurrentBalance": sum(income) - sum(expense),
        }

    def view_summary(self):
        data = self.get_summary_data()
        print("----------SUMMARY----------")
        print(f"Income Total: ${data['IncomeTotal']} - Income Count: {data['IncomeCount']}")
        print(f"Expense Total: ${data['ExpenseTotal']} - Expense Count: {data['ExpenseCount']}")
        print(f'Current Balance: ${data["CurrentBalance"]}')
        print("----------END SUMMARY----------")

    def view_transactions(self):
        print("  \tAmount | Processed | Type | Category | Description")
        for i, tran in enumerate(self.transactions):
            counter = i + 1
            print(f"{counter}:\t${tran['amount']} | "
                  f"{tran['created']} | "
                  f"{TransactionType(tran['transaction_type']).name} | "
                  f"{tran['category']} | "
                  f"{tran['description']}")


class PynanceVis:

    def __init__(self, tracker: Pynance):
        self.tracker = tracker

    def expense_by_category(self):
        # create engine and query
        eng = create_engine(self.tracker.db_file)
        query = f"SELECT * FROM transactions WHERE transaction_type = {TransactionType.EXPENSE.value}"
        # use pandas to query db for expense transactions
        df = pd.read_sql_query(query, eng)
        # group by category, sum amount
        grouped = df[['category', 'amount']].groupby('category').sum().reset_index()
        # break the amounts and the categories out
        amounts = grouped['amount']
        categories = grouped['category']
        # create visualization
        _, ax = plt.subplots()
        ax.pie(amounts, labels=categories, autopct='%1.1f%%', wedgeprops=dict(width=0.5))
        plt.title("Expense Transactions by Category")
        plt.show()

    def income_vs_expense_by_month(self):
        eng = create_engine(self.tracker.db_file)
        query = f"SELECT * FROM transactions WHERE transaction_type in (0, 1);"
        df = pd.read_sql_query(query, eng)
        df['created'] = pd.to_datetime(df['created']).dt.strftime('%Y-%m')
        income = df[df['transaction_type'] == TransactionType.INCOME.value][
            ['created', 'amount']
        ].groupby('created').sum().reset_index()
        expense = df[df['transaction_type'] == TransactionType.EXPENSE.value][
            ['created', 'amount']
        ].groupby('created').sum().reset_index()

        months = income['created']
        income_data = income['amount']
        expense_data = expense['amount']

        x = np.arange(len(months))

        width = 0.35

        _, ax = plt.subplots()

        rects1 = ax.bar(x - width / 2, income_data, width, label='Income', color='green')
        rects2 = ax.bar(x + width / 2, expense_data, width, label='Expenses', color='red')
        ax.set_xticks(x)
        ax.set_xticklabels(months)
        ax.set_ylabel('Amount')
        ax.set_title('Income vs Expenses by Month')

        ax.legend()

        plt.show()


class PynanceReport:

    def __init__(self, tracker: Pynance, title='TransactionsReport'):
        self.tracker = tracker
        eng = create_engine(self.tracker.db_file)
        today = datetime.now()
        three_months_back = (today - relativedelta(months=3)).replace(
            day=1, hour=0, minute=0, second=0, microsecond=0
        ).strftime(DATE_FORMAT)
        start_date = three_months_back.split(' ')[0]
        query = f"SELECT * FROM transactions WHERE transaction_type in (0, 1) and date([created]) >= date({start_date});"
        self.df = pd.read_sql_query(query, eng)
        self.df['created'] = pd.to_datetime(self.df['created'])
        end_date = max(self.df['created']).strftime(DATE_FORMAT).split(' ')[0]
        self.xlsx_path = Path(Path().home() / f'{title}_{start_date.replace('-','')}_{end_date.replace('-','')}.xlsx')
        self.workbook = Workbook()

    def __df_shape_to_xl_range(self, df: pd.DataFrame, start_col: str = 'A',
                               start_row: int = 1):
        end_col = chr(ord('A') + (len(df.columns) - 1))
        end_row = df.shape[0] + start_row
        return f'{start_col}{start_row}:{end_col}{end_row}'

    def __convert_range_to_table(self, xl_range: str, table_name: str, ws: Worksheet):
        table = Table(displayName=table_name, ref=xl_range)
        style = TableStyleInfo(
            name='TableStyleMedium9',
            showRowStripes=True
        )
        table.tableStyleInfo = style
        ws.add_table(table)

    def __df_to_ws(self, df: pd.DataFrame, ws: Worksheet):
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

    def create_report(self):

        ws = self.workbook.active

        ws.title = 'Transactions'

        self.__df_to_ws(self.df, ws)

        table_ref = self.__df_shape_to_xl_range(self.df)
        self.__convert_range_to_table(table_ref, 'Transactions', ws)

        self.workbook.create_sheet(title="ExpensesGrouped")
        group_ws = self.workbook['ExpensesGrouped']

        expenses = self.df[self.df['transaction_type'] == TransactionType.EXPENSE.value]
        expenses['month'] = expenses['created'].dt.to_period('M')
        grouped_df = expenses.groupby(['month', 'category']).agg({'amount': 'sum'}).reset_index()
        grouped_df['month'] = grouped_df['month'].dt.to_timestamp()

        self.__df_to_ws(grouped_df, group_ws)
        self.__convert_range_to_table(self.__df_shape_to_xl_range(grouped_df),
                                      'ExpensesByCategory',
                                      group_ws)

        self.workbook.save(str(self.xlsx_path))
        self.workbook.close()
        print(f'report saved to file path: {str(self.xlsx_path)}')